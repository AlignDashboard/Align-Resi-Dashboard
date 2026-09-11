#!/usr/bin/env python3
"""Checks for parse_daily_leasing and parse_renewal_tracker.

Workbooks are built in a temp dir rather than read from fixtures: the real
files carry resident names, so they are gitignored and CI has never seen one.
Every builder below reproduces a shape observed in a real export, and the
comment on each check names the failure it is there to catch.

Run: python scripts/test_leasing_and_renewal.py
"""
import os
import shutil
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_daily_leasing import parse as parse_leasing              # noqa: E402
from parse_renewal_tracker import parse as parse_renewal            # noqa: E402
from xlsx_anchors import LayoutError                                # noqa: E402

PASS = FAIL = 0
TMP = tempfile.mkdtemp(prefix="align-leasing-")


def check(name, ok, note=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"  [ok]   {name}")
    else:
        FAIL += 1
        print(f"  [FAIL] {name}" + (f" — {note}" if note else ""))


LEASE_HDR = ["APT #", "MKT (M) / AHP (A)", "FLOOR PLAN", "BRXBA", "SIZE (SQFT)",
             "LEASE RENT", "GROSS $/SQFT", "TOTAL RENT CONCESSION",
             "NET RENT VALUE", "NET $/SQFT", "PRIOR LEASE RATE", "TRADE OUT $",
             "TRADE OUT %", "SCHEDULED MI DATE", "LEASE TERM",
             "LEASING ASSOCIATE"]


def lease_row(unit, rent, prior, sqft=600, to_amt=None, to_pct=None):
    """A NEW LEASES row. to_amt/to_pct default to the honest arithmetic."""
    amt = (rent - prior) if to_amt is None else to_amt
    pct = ((rent - prior) / prior) if to_pct is None else to_pct
    return [unit, "MKT", "laa1", "1x1", sqft, rent, rent / sqft, 0, rent,
            rent / sqft, prior, amt, pct, "2026-09-06", 12, "Edwin"]


def build_leasing(path, leases, *, sheet_property="The Landing", units=263,
                  week_label="Week Ending", week_value="2026-09-06",
                  info_property="The Landing", trailer=True, sheet="Weekly_Leases"):
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Information"
    info["A2"], info["B2"] = "Property", info_property
    info["A4"], info["B4"] = "Units", units

    ws = wb.create_sheet(sheet)
    ws["B1"] = "Daily Activity Report"
    ws["B5"], ws["E5"] = "Property:", sheet_property
    ws["B7"], ws["E7"] = "Units:", units
    ws["B9"], ws["F9"], ws["G9"] = "Daily", week_label, week_value
    for i, h in enumerate(LEASE_HDR):
        ws.cell(row=25, column=2 + i, value=h)
    r = 26
    for lease in leases:
        for i, v in enumerate(lease):
            ws.cell(row=r, column=2 + i, value=v)
        r += 1
    if trailer:
        # The two blocks that follow the leases, laid out as the real file lays
        # them out: their columns overlap the lease table's, so a cancellation
        # row has a unit in the lease table's APT # column and a value in its
        # LEASE RENT column (the cancel block's SCHEDULED MI DATE sits there).
        # That overlap is why the STOP patterns exist.
        ws.cell(row=r + 1, column=2, value="WEEKLY AVERAGE")
        ws.cell(row=r + 1, column=6, value=621)
        ws.cell(row=r + 1, column=7, value=99999)
        ws.cell(row=r + 3, column=2, value="CANCEL/DENIALS")
        for i, h in enumerate(["APT #", "BLDG", "CANCEL DENIAL WAITLIST",
                               "HOH LAST NAME", "LEASE RENT", "SCHEDULED MI DATE",
                               "LEASE TERM", "REASON"]):
            ws.cell(row=r + 4, column=2 + i, value=h)
        ws.cell(row=r + 5, column=2, value=901)
        ws.cell(row=r + 5, column=6, value=7777)
        ws.cell(row=r + 5, column=7, value="2026-10-01")
    wb.save(path)
    return path


print("parse_daily_leasing")

p = build_leasing(f"{TMP}/a.xlsx", [lease_row(650, 5973, 3351, 621)])
r = parse_leasing(p)
check("reads the NEW LEASES block", r["totals"]["leases"] == 1
      and r["leases"][0]["prior_rent"] == 3351,
      f"got {r['totals']}")
check("carries the prior rent, which nothing else in the pipeline has",
      r["leases"][0]["prior_rent"] == 3351)
check("week-ending read from a real date cell", r["as_of"] == "2026-09-06",
      r["as_of"])

# Two separate guards keep the trailing sections out, and each is checked on
# its own.
#
# STOP is the primary one. The WEEKLY AVERAGE row is the first marker after the
# leases and it carries a NUMBER in the lease-rent column (the week's average
# rent), so the type guard below cannot reject it -- only the marker can. Drop
# the STOP patterns and this week reports two leases, one of them an average.
stop_check = [c for c in r["checks"] if c["check"].startswith("stopped")][0]
check("halts at the first trailing section marker", stop_check["ok"]
      and "weekly average" in stop_check["note"], stop_check)
check("the numeric WEEKLY AVERAGE row is not counted as a lease",
      r["totals"]["leases"] == 1
      and all(str(l["unit"]).isdigit() for l in r["leases"]),
      [l["unit"] for l in r["leases"]])

# The type guard is the backstop for anything STOP does not catch: the
# cancellation block puts its scheduled move-in DATE where a lease's rent goes,
# and left unguarded the totals raise TypeError on str + int, taking the whole
# pipeline pass down rather than reporting one bad row.
check("totals stay numeric, so a stray date cannot crash the parse",
      isinstance(r["totals"]["lease_rent"], (int, float)), r["totals"])

p = build_leasing(f"{TMP}/b.xlsx", [lease_row(1, 5000, 4000), lease_row(2, 6000, 5000)])
r = parse_leasing(p)
check("totals sum the week", r["totals"]["lease_rent"] == 11000
      and r["totals"]["prior_rent"] == 9000, r["totals"])

# A row whose reported trade-out disagrees with its own rent and prior rate is
# the signal that a column moved. Without the arithmetic check it publishes.
p = build_leasing(f"{TMP}/c.xlsx", [lease_row(3, 5000, 4000, to_amt=2500)])
r = parse_leasing(p)
check("a lease that fails its own trade-out arithmetic is reported",
      any("does not reconcile" in x for x in r["problems"])
      and not [c for c in r["checks"] if c["check"].startswith("every lease")][0]["ok"],
      r["problems"])

# Chorus's real file: the Information sheet still says The Landing.
p = build_leasing(f"{TMP}/Chorus - Daily Report.xlsx", [lease_row(1901, 4654, 4047)],
                  sheet_property="Chorus", info_property="The Landing", units=416)
r = parse_leasing(p)
check("the filename outranks both in-file sources", r["property"] == "Chorus",
      r["property"])
check("a property disagreement is reported, not silently resolved",
      any("disagrees with itself" in x for x in r["problems"]), r["problems"])

# The Landing's own file names no property at all, so the sheet header is what
# resolves it -- the Information sheet alone would be a coin toss.
p = build_leasing(f"{TMP}/Daily Report- Week Ending 9.7.26.xlsx",
                  [lease_row(650, 5973, 3351)], sheet_property="The Landing")
r = parse_leasing(p)
check("a filename with no property falls back to the sheet header",
      r["property"] == "The Landing" and r["property_from_filename"] is None,
      r["property"])

# Chorus writes "Week To Date" beside the text "Ending 09.13.26".
p = build_leasing(f"{TMP}/d.xlsx", [lease_row(1, 5000, 4000)],
                  week_label="Week To Date", week_value="Ending 09.13.26")
r = parse_leasing(p)
check("week-ending read from the M.D.YY text form", r["as_of"] == "2026-09-13",
      r["as_of"])

p = build_leasing(f"{TMP}/Report Week Ending 9.13.26.xlsx",
                  [lease_row(1, 5000, 4000)], week_label="x", week_value="x")
r = parse_leasing(p)
check("week-ending falls back to the filename", r["as_of"] == "2026-09-13",
      r["as_of"])

# The week LABEL, which the store keys on. The sheet's own cell drifts between
# snapshots of one week -- the 2026-08-31 and 2026-09-08 copies of the real
# "Week Ending 9.7.26" say 2026-09-07 and 2026-09-06 -- so keying on it files
# one week under two keys and counts it twice.
a = build_leasing(f"{TMP}/2026-08-31 Daily Report- Week Ending 9.7.26.xlsx",
                  [lease_row(1, 5000, 4000)], week_value="2026-09-07")
b = build_leasing(f"{TMP}/2026-09-08 Daily Report- Week Ending 9.7.26.xlsx",
                  [lease_row(1, 5000, 4000), lease_row(2, 6000, 5000)],
                  week_value="2026-09-06")
ra, rb = parse_leasing(a), parse_leasing(b)
check("two snapshots of one week get ONE week label",
      ra["as_of"] == rb["as_of"] == "2026-09-07",
      f"{ra['as_of']} vs {rb['as_of']}")
check("the sheet's own answer is kept beside it",
      ra["week_ending_sheet"] == "2026-09-07" and rb["week_ending_sheet"] == "2026-09-06",
      f"{ra['week_ending_sheet']} / {rb['week_ending_sheet']}")
check("a drifting week cell is reported, not silently resolved",
      any("the sheet says" in x for x in rb["problems"]) and not ra["problems"],
      rb["problems"])

# Chorus names a RANGE; its week is the end of it, not the start.
p = build_leasing(f"{TMP}/2026-09-09 09.07.2026- 09.13.2026- Chorus - Daily Report.xlsx",
                  [lease_row(1, 5000, 4000)], sheet_property="Chorus",
                  week_label="Week To Date", week_value="Ending 09.13.26")
r = parse_leasing(p)
check("a range-style filename takes the END of the range",
      r["as_of"] == "2026-09-13", r["as_of"])

# A name carrying only the filer's arrival prefix says nothing about the week.
p = build_leasing(f"{TMP}/2026-09-08 Weekly Report.xlsx",
                  [lease_row(1, 5000, 4000)], week_value="2026-09-06")
r = parse_leasing(p)
check("a filename with only the arrival prefix falls back to the sheet",
      r["as_of"] == "2026-09-06", r["as_of"])

p = build_leasing(f"{TMP}/e.xlsx", [lease_row(1, 5000, 4000)], sheet="Something_Else")
try:
    parse_leasing(p)
    check("a workbook with no Weekly_Leases sheet is refused", False, "parsed anyway")
except LayoutError:
    check("a workbook with no Weekly_Leases sheet is refused", True)


# --------------------------------------------------------------------------
print("\nparse_renewal_tracker")

MODERN = ["STATUS", "NAMES", "UNIT", "CURRENT LEASE END DATE", "FLOOR PLAN",
          "UNIT TYPE", "NOTES", "CURRENT GROSS RENT", "CURRENT TERM",
          "MARKET RATE", "Current LTL", "Current LTL %", "BEST OFFER RATE",
          "BEST OFFER DIFFERENCE", "BEST OFFER %"]
# June 2025's shape: no offered-rent column, BEST OFFER $ IS the rate
JUNE25 = ["STATUS", "CURRENT LEASE END DATE", "FLOOR PLAN", "UNIT TYPE", "UNIT",
          "NOTES", "CURRENT RENT", "MARKET RATE", "LTL", "LTL %",
          "CURRENT TERM", "BEST OFFER %", "BEST OFFER $"]
# May 2025's shape: BEST OFFER $ is the DIFFERENCE and OFFER RATE is the rate
MAY25 = ["STATUS", "CURRENT LEASE END DATE", "FLOOR PLAN", "NAME", "UNIT",
         "NOTES", "CURRENT EFFECTIVE RENT", "MARKET RATE", "LTL", "LTL %",
         "CURRENT TERM", "BEST OFFER %", "BEST OFFER $", "OFFER RATE"]

MTM_HDR = ["STATUS", "Unit Type", "Unit", "Code", "Name",
           "Month to Month Charge", "MTM Charge as % of Actual Rent",
           "Lease From", "Lease To", "Current Rent", "Market Rent"]


def build_renewal(path, sheets, mtm_rows=None, header_row=9):
    """sheets: {title: (header list, [row lists])}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, (hdr, rows) in sheets.items():
        ws = wb.create_sheet(title)
        for i, h in enumerate(hdr):
            ws.cell(row=header_row, column=1 + i, value=h)
        for j, row in enumerate(rows):
            for i, v in enumerate(row):
                ws.cell(row=header_row + 1 + j, column=1 + i, value=v)
    if mtm_rows is not None:
        ws = wb.create_sheet("MTM")
        for i, h in enumerate(MTM_HDR):
            ws.cell(row=9, column=1 + i, value=h)
        for j, row in enumerate(mtm_rows):
            for i, v in enumerate(row):
                ws.cell(row=10 + j, column=1 + i, value=v)
    wb.save(path)
    return path


# modern: BEST OFFER RATE is the rate, BEST OFFER DIFFERENCE the difference
modern_rows = [["Renewed", "A Name", 836, "Sep-03-2026", "2B", "lab19", "", 4791,
                12, 7047, -2256, 0.32, 5222, 431, 0.15]]
p = build_renewal(f"{TMP}/Landing 2025 Renewal Tracker.xlsx",
                  {"September 2026": (MODERN, modern_rows)},
                  mtm_rows=[["Waiting", "lab19", 355, "t0519674", "A Name",
                             None, None, "2024-02-26", "2025-02-25", 4109, 7882]])
r = parse_renewal(p)
m = r["months"][0]
check("month sheet found from its title", m["month"] == "2026-09", m["month"])
check("offered rent taken as a RATE where the column is one",
      m["offered_rent"] == 5222 and m["current_rent"] == 4791, m)
# The sheet's own BEST OFFER % says 0.15 here while the rents say 9.0%; the
# published increase must come from the rents. (1e-6, not exact: the parser
# rounds to six places like every other ratio in the pipeline.)
check("weighted increase computed from the totals, not the sheet's own %",
      abs(m["wtd_increase"] - (5222 / 4791 - 1)) < 1e-6
      and abs(m["wtd_increase"] - 0.15) > 0.05, m["wtd_increase"])
check("property resolved from the filename", r["property"] == "The Landing",
      r["property"])
check("MTM roster read", r["mtm"]["units"] == 1 and r["mtm"]["current_rent"] == 4109,
      r["mtm"])
# The tenant code must reach store_report under a name the scrub knows.
check("the MTM tenant code is emitted as resident_code, so scrub() drops it",
      "resident_code" in r["mtm"]["rows"][0], sorted(r["mtm"]["rows"][0]))

# THE key guard. June 2025's "BEST OFFER $" is the offered RATE; May 2025's is
# the DIFFERENCE. Resolving by label instead of by arithmetic reads 4326+4542
# = 8868 as the offer -- a 105% increase that would publish as real.
june_rows = [["Notice", "Jun-03-2025", "2B", "2B1B", 433, "", "$4,326", "$5,232",
              -906, -0.2, 13, "4.99%", "$4,542"]]
p = build_renewal(f"{TMP}/Landing tracker june.xlsx",
                  {"June 2025": (JUNE25, june_rows)}, header_row=10)
m = parse_renewal(p)["months"][0]
check("a rate-shaped 'BEST OFFER $' is read as the rate",
      m["offered_rent"] == 4542, m)
check("that month's increase is ~5%, not ~105%",
      m["wtd_increase"] is not None and 0.04 < m["wtd_increase"] < 0.06,
      m["wtd_increase"])
check("money written as '$4,326' text is parsed, not dropped",
      m["current_rent"] == 4326, m["current_rent"])

# OFFER RATE is " ha" on the real sheet's first row -- stray text in a numeric
# column -- so the offered rent can only come from BEST OFFER $, the DIFFERENCE.
may_rows = [["Notice", "May-30-2025", "lac2a", "A Name", 223, "", 4932, 6085,
             -1153, -0.23, 12, 0.099, 488.268, " ha"]]
p = build_renewal(f"{TMP}/Landing tracker may.xlsx",
                  {"May 2025": (MAY25, may_rows)}, header_row=10)
m = parse_renewal(p)["months"][0]
check("a difference-shaped column is added to the current rent",
      m["offered_rent"] is not None and abs(m["offered_rent"] - 5420.268) < 1,
      m["offered_rent"])
check("the difference path is what resolved it, not a rate column",
      any("difference" in b for b in m["offer_basis"]), m["offer_basis"])
check("that month's increase is ~9.9%, not ~10%+ of a doubled rent",
      0.09 < m["wtd_increase"] < 0.11, m["wtd_increase"])

# A sheet whose header cannot be matched must be recorded, never guessed at.
p = build_renewal(f"{TMP}/Landing tracker mixed.xlsx",
                  {"September 2026": (MODERN, modern_rows),
                   "March 2024": (["Colour", "Shape", "Notes"], [["red", "square", ""]])})
r = parse_renewal(p)
check("an unreadable month sheet is recorded, not guessed at",
      any(u["sheet"] == "March 2024" for u in r["unread_sheets"])
      and len(r["months"]) == 1, r["unread_sheets"])
check("skipped sheets are surfaced as a problem",
      any("could not be matched" in x for x in r["problems"]), r["problems"])

# Two sheets can name one month ("July 2025" and "July 2025 (2)").
p = build_renewal(f"{TMP}/Landing tracker dup.xlsx",
                  {"July 2025": (MODERN, modern_rows * 3),
                   "July 2025 (2)": (MODERN, modern_rows)})
r = parse_renewal(p)
check("a month on two sheets keeps the fuller one and says so",
      len(r["months"]) == 1 and r["months"][0]["leases"] == 3
      and any("two sheets" in x for x in r["problems"]), r["problems"])

p = build_renewal(f"{TMP}/tracker with no property.xlsx",
                  {"September 2026": (MODERN, modern_rows)})
r = parse_renewal(p)
check("a filename naming no property is reported",
      r["property"] is None and any("names no known property" in x
                                    for x in r["problems"]), r["problems"])

try:
    p = build_renewal(f"{TMP}/Landing tracker empty.xlsx",
                      {"Summary": (["a", "b", "c"], [[1, 2, 3]])})
    parse_renewal(p)
    check("a workbook with no readable month sheet is refused", False, "parsed anyway")
except LayoutError:
    check("a workbook with no readable month sheet is refused", True)

shutil.rmtree(TMP, ignore_errors=True)
print(f"\n{PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
