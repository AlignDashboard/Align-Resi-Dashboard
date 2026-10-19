#!/usr/bin/env python3
"""Guard tests: the build reuses a parse only when nothing it depends on has
changed, never keeps a person between runs, and checks itself every run.

parse_cache.py lets process_manifest skip the one step that made the build
grow with the drop tree -- re-parsing every report Drive has ever held. A
cache that is wrong is worse than none: it would serve a stale parse with
every step green. So this holds down, without a fixture or the network:

  1. the key -- reused when the file, the parser, the helpers it imports and
     the config it names are all unchanged; read again when any one moves;
     NOT read again for a change to a file the parser never reads;
  2. what may be kept -- a type off the list is never kept, a parse that
     carries a person is never kept, a failure is kept as its message, and a
     failure of the machine is not;
  3. the check -- one kept parse per parser re-read and compared each run,
     least recently checked first, and a parse whose key misses a dependency
     sets that parser's whole cache aside;
  4. housekeeping -- unreadable entries, a lost index, pruning, PARSE_CACHE=off;
  5. the real parsers' keys -- what they import and which config they name;
  6. end to end through process_manifest with the real daily leasing parser:
     a warm run stores exactly what a run with no cache stores, parses
     nothing, and leaves no name anywhere in the cache directory;
  7. the fetch, through fetch_drive.main and process_manifest against a
     stubbed Drive: a file whose parse is kept is left in Drive and keyed on
     Drive's own md5, the build stores exactly what it stores with every
     file on disk, a file it needs after all is downloaded then -- and a
     file it cannot get ends the run rather than going missing from it;
  8. the contract around it: nothing the daily run executes reads a file the
     fetch may leave in Drive, and update.yml restores the cache first.

Run: python scripts/test_parse_cache.py
"""
import contextlib
import datetime as dt
import importlib
import io
import json
import os
import pathlib
import shutil
import sys
import tempfile
import warnings

sys.dont_write_bytecode = True    # a rewritten stub must never run from stale bytecode

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import build_metrics as bm        # noqa: E402
import parse_cache as pc          # noqa: E402

PASS = FAIL = 0
T0 = dt.datetime(2026, 10, 6, 12, 0, tzinfo=dt.timezone.utc)


def ok(name, cond, detail=None):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}" + (f"\n        {detail}" if detail is not None else ""))


# A parser the tests rewrite. It records every call, imports a helper, reads a
# config file by name, and reads one thing its key cannot see -- an environment
# variable -- which is what the self-check has to catch.
PARSER = '''"""A stub parser.

Usage, as the real parsers document it: from quoted_only import nothing
"""
import json
import os

import helper

HERE = os.path.dirname(os.path.abspath(__file__))
VERSION = 1


def parse(path):
    with open(os.path.join(HERE, "calls.log"), "a") as fh:
        fh.write(os.path.basename(path) + "\\n")
    cfg = json.load(open(os.path.join(HERE, "..", "config", "props.json")))
    text = open(path).read()
    if text.startswith("FAIL"):
        raise ValueError(f"{os.path.basename(path)}: no 'Weekly_Leases' sheet")
    if text.startswith("OSERR"):
        raise OSError("too many open files")
    out = {"body": text, "helper": helper.TAG, "config": cfg["tag"],
           "version": VERSION, "rows": [{"unit": 1, "rent": 5000.5}]}
    if text.startswith("PERSON"):
        out["rows"][0]["resident_name"] = "Resident Zed"
    if os.environ.get("STUB_PARSER_EXTRA"):
        out["extra"] = os.environ["STUB_PARSER_EXTRA"]
    return out
'''


class Sandbox:
    def __init__(self, tmp):
        self.root = pathlib.Path(tmp)
        self.scripts = self.root / "scripts"
        self.config = self.root / "config"
        self.files = self.root / "files"
        self.cache = self.root / "cache"
        for d in (self.scripts, self.config, self.files):
            d.mkdir(parents=True)
        (self.scripts / "stub_parser.py").write_text(PARSER)
        (self.scripts / "helper.py").write_text("TAG = 'helper-1'\n")
        (self.scripts / "quoted_only.py").write_text("nothing = 1\n")
        (self.config / "props.json").write_text(json.dumps({"tag": "props-1"}))
        (self.config / "report_map.json").write_text(json.dumps({"tag": "map-1"}))
        sys.path.insert(0, str(self.scripts))
        self.hour = 0

    def item(self, name, text=None, rtype="stub"):
        p = self.files / name
        if text is not None:
            p.write_text(text)
        return {"report_type": rtype, "parser": "stub_parser", "path": str(p),
                "name": name}

    def calls(self):
        log = self.scripts / "calls.log"
        got = log.read_text().split() if log.exists() else []
        log.unlink(missing_ok=True)
        return got

    def run(self, items, types=None, hours=1, touch=False):
        """One build's worth: a fresh ParseCache over the same directory.

        touch=True writes into each parse between parse() and finish(), as
        process_manifest does with landed_at and source_file."""
        self.hour += hours
        logs = []
        cache = pc.ParseCache(self.cache, person_fields=bm.PII_FIELDS,
                              cached_types={"stub": "test"} if types is None else types,
                              config_dir=self.config,
                              now=T0 + dt.timedelta(hours=self.hour), log=logs.append)
        cache.prepare(items, load)
        out = []
        for it in items:
            try:
                got = cache.parse(it, load(it))
                if touch:
                    got["landed_at"] = "written by the caller"
                out.append(got)
            except Exception as e:                      # noqa: BLE001
                out.append((type(e).__name__, str(e)))
        cache.finish()
        return out, logs, cache


def load(item):
    import helper                                       # noqa: F401
    importlib.reload(sys.modules["helper"])
    mod = importlib.import_module(item["parser"])
    return importlib.reload(mod)


def section_key(sb):
    print("1. the key")
    a, b, c = sb.item("a.txt", "alpha"), sb.item("b.txt", "bravo"), sb.item("c.txt", "charlie")
    first, _, cache = sb.run([a, b, c], touch=True)
    ok("a cold run parses every file", sorted(sb.calls()) == ["a.txt", "b.txt", "c.txt"]
       and cache.counts["parsed"] == 3)
    again, logs, cache = sb.run([a, b, c])
    called = sb.calls()
    ok("a warm run reuses every parse, re-reading exactly one to check it",
       len(called) == 1 and cache.counts["reused"] == 3 and cache.counts["parsed"] == 0,
       (called, cache.counts))
    ok("the reused parse is the parser's own, not what the caller did to it",
       again[0] == {"body": "alpha", "helper": "helper-1", "config": "props-1",
                    "version": 1, "rows": [{"unit": 1, "rent": 5000.5}]}, again[0])

    sb.item("b.txt", "bravo, edited in place")
    out, _, cache = sb.run([a, b, c])
    sb.calls()
    ok("changed bytes are parsed again, and only they are",
       (cache.counts["parsed"], cache.counts["reused"]) == (1, 2)
       and out[1]["body"] == "bravo, edited in place", cache.counts)

    # Counted, not inferred from the call log: the self-check re-reads one
    # reused file per run, so "the parser was called" proves nothing.
    moved = sb.item("renamed.txt", "alpha")
    _, logs, cache = sb.run([moved])
    sb.calls()
    ok("the same bytes under another name are parsed again (parsers read names)",
       (cache.counts["parsed"], cache.counts["reused"]) == (1, 0)
       and not any("differs" in l for l in logs), cache.counts)

    src = (sb.scripts / "stub_parser.py")
    src.write_text(src.read_text().replace("VERSION = 1", "VERSION = 2"))
    out, _, cache = sb.run([a, b, c])
    sb.calls()
    ok("a change to the parser's own code parses everything again",
       cache.counts["parsed"] == 3 and out[0]["version"] == 2, cache.counts)

    (sb.scripts / "helper.py").write_text("TAG = 'helper-2'\n")
    out, _, cache = sb.run([a, b, c])
    sb.calls()
    ok("a change to a module the parser imports parses everything again",
       cache.counts["parsed"] == 3 and out[0]["helper"] == "helper-2", cache.counts)

    (sb.config / "props.json").write_text(json.dumps({"tag": "props-2"}))
    out, _, cache = sb.run([a, b, c])
    sb.calls()
    ok("a change to a config file the parser names parses everything again",
       cache.counts["parsed"] == 3 and out[0]["config"] == "props-2", cache.counts)

    (sb.config / "report_map.json").write_text(json.dumps({"tag": "map-2"}))
    _, _, cache = sb.run([a, b, c])
    sb.calls()
    ok("a config file the parser never names does not (report_map.json)",
       (cache.counts["parsed"], cache.counts["reused"]) == (0, 3), cache.counts)

    (sb.scripts / "quoted_only.py").write_text("nothing = 2\n")
    _, _, cache = sb.run([a, b, c])
    sb.calls()
    ok("nor does a module its docstring merely quotes -- imports come from ast",
       (cache.counts["parsed"], cache.counts["reused"]) == (0, 3), cache.counts)


def section_kept(sb):
    print("\n2. what may be kept")
    x = sb.item("x.txt", "xray", rtype="unlisted")
    sb.run([x])
    _, _, cache = sb.run([x])
    sb.calls()
    ok("a report type off the list is parsed on every run",
       (cache.counts["parsed"], cache.counts["reused"]) == (1, 0), cache.counts)

    p = sb.item("p.txt", "PERSON row")
    _, logs, _ = sb.run([p])
    sb.calls()
    out, _, cache = sb.run([p])
    ok("a parse that carries a person is never kept, whatever its type",
       sb.calls() == ["p.txt"] and cache.counts["reused"] == 0)
    ok("...and the run says so", any("person field ('resident_name')" in l for l in logs), logs)
    held = b"".join(f.read_bytes() for f in sb.cache.rglob("*") if f.is_file())
    ok("...and the name is nowhere in the cache directory", b"Resident Zed" not in held)

    f = sb.item("f.txt", "FAIL here", rtype="unlisted")
    first, _, _ = sb.run([f])
    sb.calls()
    again, _, cache = sb.run([f])
    called = sb.calls()
    ok("a failure is kept for any type, and raised again with its own message",
       again[0] == ("ReusedFailure", first[0][1]) and first[0][0] == "ValueError"
       and cache.counts["reused_failures"] == 1, (first, again))
    ok("...without the parser being asked again, beyond the one check",
       called == ["f.txt"], called)

    o = sb.item("o.txt", "OSERR", rtype="stub")
    sb.run([o])
    _, _, cache = sb.run([o])
    sb.calls()
    ok("a failure of the machine (OSError) is not kept",
       (cache.counts["parsed"], cache.counts["reused_failures"]) == (1, 0), cache.counts)

    real = pc.CACHED_TYPES.keys() & pc.NOT_KEPT.keys()
    ok("the kept and never-kept lists are disjoint", not real, real)
    ok("rent rolls, delinquency, renewal trackers and burn-offs are never kept",
       {"rent_roll", "ar_analytics", "renewal_tracker", "concession_burnoff"}
       <= pc.NOT_KEPT.keys() and not ({"rent_roll", "ar_analytics", "renewal_tracker",
                                       "concession_burnoff"} & pc.CACHED_TYPES.keys()))
    types = {e["report_type"] for e in
             json.load(open(os.path.join(REPO, "config", "report_map.json")))["subfolders"]}
    ok("every listed type is a real report type (a typo would keep nothing)",
       set(pc.CACHED_TYPES) | set(pc.NOT_KEPT) <= types,
       (set(pc.CACHED_TYPES) | set(pc.NOT_KEPT)) - types)


def section_check(sb):
    print("\n3. the check")
    shutil.rmtree(sb.cache, ignore_errors=True)
    items = [sb.item(n, n.upper()) for n in ("k1.txt", "k2.txt", "k3.txt")]
    sb.run(items)
    sb.calls()
    seen = []
    for _ in range(3):
        _, logs, cache = sb.run(items)
        seen += [c[1] for c in cache.checks]
        ok_line = any(l.startswith("[cache] checked stub_parser") for l in logs)
    sb.calls()
    ok("each run checks one kept parse, least recently checked first",
       sorted(seen) == ["k1.txt", "k2.txt", "k3.txt"], seen)
    ok("...and logs the match", ok_line)

    os.environ["STUB_PARSER_EXTRA"] = "outside the key"
    try:
        out, logs, cache = sb.run(items)
        called = sb.calls()
        ok("a parse depending on something outside its key is caught by the check",
           any("differs from the one kept" in l for l in logs), logs)
        ok("...every kept parse of that parser is set aside and read again",
           sorted(set(called)) == ["k1.txt", "k2.txt", "k3.txt"], called)
        ok("...the checked file is not parsed twice",
           len(called) == 3, called)
        ok("...and the run gets the fresh parses",
           all(o.get("extra") == "outside the key" for o in out), out)
        out, _, cache = sb.run(items)
        ok("the replacements are kept for the next run",
           cache.counts["reused"] == 3 and all(o.get("extra") for o in out))
    finally:
        del os.environ["STUB_PARSER_EXTRA"]
        sb.calls()

    # A failure kept for a type whose successful parse may NOT be kept (a
    # renewal tracker, say) must not outlive the run that disproves it.
    shutil.rmtree(sb.cache, ignore_errors=True)
    t = sb.item("t.txt", "FAIL at first", rtype="unlisted")
    sb.run([t])
    sb.item("t.txt", "FAIL at first")              # same bytes...
    os.environ["STUB_PARSER_EXTRA"] = "fixed"      # ...but it parses now
    src = sb.scripts / "stub_parser.py"
    text = src.read_text()
    src.write_text(text.replace('if text.startswith("FAIL"):',
                                'if text.startswith("FAIL") and not os.environ.get("STUB_PARSER_EXTRA"):'))
    try:
        # the parser changed, so its key changed: an ordinary miss
        out, _, _ = sb.run([t])
        ok("a file that failed and now parses is read afresh", isinstance(out[0], dict), out)
        out, _, cache = sb.run([t])
        ok("...and its old failure is not served on the next run",
           isinstance(out[0], dict) and cache.counts["reused_failures"] == 0, out)
    finally:
        del os.environ["STUB_PARSER_EXTRA"]
        src.write_text(text)
        sb.calls()

    # The same, but with the key unchanged: only the check can find it.
    shutil.rmtree(sb.cache, ignore_errors=True)
    src.write_text(text.replace('if text.startswith("FAIL"):',
                                'if text.startswith("FAIL") and not os.environ.get("STUB_PARSER_EXTRA"):'))
    try:
        sb.run([t])                                 # kept as a failure
        os.environ["STUB_PARSER_EXTRA"] = "fixed"
        out, logs, _ = sb.run([t])                  # the check disproves it
        ok("a kept failure the check disproves is replaced by the parse",
           isinstance(out[0], dict) and any("differs" in l for l in logs), (out, logs))
        # Left behind, it would be the least recently checked entry, so every
        # later run would re-check it, disprove it again and set the whole
        # parser aside -- a warning and a full re-read on every run, forever.
        out, logs, cache = sb.run([t])
        ok("...and is gone, though the type may not keep a parse: the next run "
           "neither serves it nor sets the parser aside again",
           isinstance(out[0], dict) and cache.counts["reused_failures"] == 0
           and not any("differs" in l for l in logs), logs)
    finally:
        os.environ.pop("STUB_PARSER_EXTRA", None)
        src.write_text(text)
        sb.calls()


def section_housekeeping(sb):
    print("\n4. housekeeping")
    shutil.rmtree(sb.cache, ignore_errors=True)
    a, b = sb.item("h1.txt", "one"), sb.item("h2.txt", "two")
    sb.run([a, b])
    sb.calls()
    for f in (sb.cache / "entries").glob("*.pickle"):
        f.write_bytes(b"not a pickle")
    out, logs, _ = sb.run([a, b])
    ok("an unreadable entry is parsed afresh, and the run goes on",
       [o["body"] for o in out] == ["one", "two"] and any("could not be read" in l for l in logs))
    sb.calls()

    (sb.cache / "index.json").unlink()
    _, _, cache = sb.run([a, b])
    ok("a lost index loses no parse -- entries are found by their key",
       cache.counts["reused"] == 2, cache.counts)
    sb.calls()

    sb.run([a], hours=24 * (pc.PRUNE_DAYS + 1))
    keys = {cache.key(b, load(b))}
    left = {p.stem for p in (sb.cache / "entries").glob("*.pickle")}
    ok(f"an entry no run has used for {pc.PRUNE_DAYS} days is pruned", not (keys & left), left)
    _, _, cache = sb.run([a])
    ok("...and one that is used is kept", cache.counts["reused"] == 1)
    sb.calls()

    # A fault inside the cache's own bookkeeping costs the bookkeeping, never
    # the build: process_manifest calls prepare() and finish() unguarded.
    broken = pc.ParseCache(sb.cache, person_fields=bm.PII_FIELDS,
                           cached_types={"stub": "test"}, config_dir=sb.config,
                           now=T0 + dt.timedelta(hours=sb.hour + 1), log=lambda m: None)
    broken._check = lambda *a: 1 / 0
    broken._record = lambda *a: 1 / 0
    try:
        broken.prepare([a], load)
        got = broken.parse(a, load(a))
        broken.finish()
        survived = got["body"] == "one"
    except Exception as e:                              # noqa: BLE001
        survived = repr(e)
    ok("a fault in the check or the bookkeeping does not stop the build", survived is True,
       survived)
    sb.calls()

    os.environ["PARSE_CACHE"] = "off"
    try:
        off = pc.from_env(person_fields=bm.PII_FIELDS)
        ok("PARSE_CACHE=off parses everything afresh", isinstance(off, pc.NoCache))
        off.prepare([a], load)
        off.parse(a, load(a))
        ok("...and keeps nothing", sb.calls() == ["h1.txt"])
    finally:
        del os.environ["PARSE_CACHE"]


def section_real_keys():
    print("\n5. the real parsers' keys")
    scripts, config = pathlib.Path(HERE), pathlib.Path(REPO) / "config"

    def names(parser):
        srcs = pc.module_closure(scripts / f"{parser}.py")
        return ({p.name for p in srcs}, {p.name for p in pc.configs_named(srcs, config)})

    s, c = names("parse_budget")
    ok("the budget's key covers the T12 parser it wraps, and the COA map",
       {"parse_budget.py", "parse_t12_statement.py"} <= s and "coa_map.json" in c, (s, c))
    s, c = names("parse_comps")
    ok("the comps' key covers xlsx_anchors and the property master",
       "xlsx_anchors.py" in s and "properties.json" in c, (s, c))
    s, c = names("parse_daily_leasing")
    ok("the daily leasing key is its own code, xlsx_anchors and the property master",
       s == {"parse_daily_leasing.py", "xlsx_anchors.py"} and c == {"properties.json"}, (s, c))
    every = {n for m in (str(p.stem) for p in scripts.glob("parse_*.py")) for n in names(m)[1]}
    ok("no parser's key covers report_map.json, the file that changes most",
       "report_map.json" not in every, every)

    # Reading a tree compiles the source. Done every run, a warning in any
    # parser would be printed in every log -- parse_lease_tradeout's docstring
    # did exactly that on the first run with the cache, on Python 3.12.
    noisy = pathlib.Path(tempfile.mkdtemp(prefix="noisy_")) / "noisy_parser.py"
    noisy.write_text('def parse(path):\n    """a pattern like `\\)?%?$`"""\n')
    with warnings.catch_warnings(record=True) as seen:
        warnings.simplefilter("always")
        pc.configs_named(pc.module_closure(noisy), config)
    ok("reading a parser's tree prints none of its compile warnings",
       not seen, [str(w.message) for w in seen])
    dirty = []
    for src in sorted({s for p in scripts.glob("parse_*.py") for s in pc.module_closure(p)}):
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            try:
                compile(src.read_text(encoding="utf-8"), str(src), "exec")
            except (SyntaxError, SyntaxWarning, DeprecationWarning) as e:
                dirty.append(f"{src.name}: {e}")
    ok("every module a parser's key covers compiles without a warning "
       "(an invalid escape is a SyntaxError in a later Python)", not dirty, dirty)


def build_week(path, rent, prior):
    """The Weekly_Leases sheet in the real layout, one lease, an associate named."""
    import openpyxl
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "Information"
    info["A2"], info["B2"] = "Property", "The Landing"
    ws = wb.create_sheet("Weekly_Leases")
    ws["B5"], ws["E5"] = "Property:", "The Landing"
    ws["B7"], ws["E7"] = "Units:", 263
    hdr = ["APT #", "MKT (M) / AHP (A)", "FLOOR PLAN", "BRXBA", "SIZE (SQFT)",
           "LEASE RENT", "GROSS $/SQFT", "TOTAL RENT CONCESSION", "NET RENT VALUE",
           "NET $/SQFT", "PRIOR LEASE RATE", "TRADE OUT $", "TRADE OUT %",
           "SCHEDULED MI DATE", "LEASE TERM", "LEASING ASSOCIATE"]
    for i, h in enumerate(hdr):
        ws.cell(row=25, column=2 + i, value=h)
    row = [901, "MKT", "laa1", "1x1", 621, rent, rent / 621, 0, rent, rent / 621,
           prior, rent - prior, (rent - prior) / prior, "2026-09-06", 12, "Edwin"]
    for i, v in enumerate(row):
        ws.cell(row=26, column=2 + i, value=v)
    ws.cell(row=28, column=2, value="WEEKLY AVERAGE")
    wb.save(path)


ROLL_STUB = '''
import json
def parse(path):
    return json.load(open(path))
'''


def section_end_to_end(tmp):
    print("\n6. end to end, through process_manifest")
    stubs = pathlib.Path(tmp) / "e2e_stubs"
    stubs.mkdir()
    (stubs / "stub_roll_parser.py").write_text(ROLL_STUB)
    sys.path.insert(0, str(stubs))
    weeks = [("2026-09-08 Daily Report- Week Ending 9.7.26.xlsx", 5973, 3351),
             ("2026-09-15 Daily Report- Week Ending 9.14.26.xlsx", 6100, 5800),
             ("2026-09-22 Daily Report- Week Ending 9.21.26.xlsx", 5400, 5500)]
    roll = {"report_type": "rent_roll", "property": "The Landing",
            "property_code": "p0005611", "as_of": "2026-09-21",
            "source_file": "RentRoll09_21_2026.xlsx", "totals": {"units": 1},
            "units": [{"unit": "901", "resident_name": "Resident Zed",
                       "rent": 5973, "occupied": True}], "checks": []}

    def build(work, landed="2026-10-05T12:00:00Z"):
        dl = work / "_downloads"
        dl.mkdir(parents=True, exist_ok=True)
        if not (work / "config").exists():
            os.symlink(os.path.join(REPO, "config"), work / "config")
        manifest = []
        for name, rent, prior in weeks:
            if not (dl / name).exists():
                build_week(dl / name, rent, prior)
            manifest.append({"report_type": "daily_leasing_report",
                             "parser": "parse_daily_leasing", "path": str(dl / name),
                             "name": name, "landed_at": landed})
        (dl / "RentRoll09_21_2026.xlsx").write_text(json.dumps(roll))
        manifest.append({"report_type": "rent_roll", "parser": "stub_roll_parser",
                         "path": str(dl / "RentRoll09_21_2026.xlsx"),
                         "name": "RentRoll09_21_2026.xlsx", "landed_at": landed})
        (dl / "manifest.json").write_text(json.dumps(manifest))

    def run(work, cache_env=None):
        shutil.rmtree(work / "data", ignore_errors=True)
        prev_cwd, prev_data = os.getcwd(), bm.DATA
        prev_env = os.environ.pop("PARSE_CACHE", None)
        if cache_env is not None:
            os.environ["PARSE_CACHE"] = cache_env
        bm.DATA = work / "data"
        out = io.StringIO()
        try:
            os.chdir(work)
            with contextlib.redirect_stdout(out):
                bm.process_manifest()
        finally:
            os.chdir(prev_cwd)
            bm.DATA = prev_data
            os.environ.pop("PARSE_CACHE", None)
            if prev_env is not None:
                os.environ["PARSE_CACHE"] = prev_env
        files = {str(p.relative_to(work / "data")): p.read_bytes()
                 for p in sorted((work / "data").rglob("*.json"))}
        return out.getvalue(), files

    base = pathlib.Path(tmp) / "e2e"
    plain, warm = base / "plain", base / "cached"
    build(plain)
    build(warm)
    log_off, want = run(plain, cache_env="off")
    ok("with the cache off, nothing is written under _cache",
       not (plain / "_cache").exists())
    _, cold = run(warm)
    log_warm, got = run(warm)
    ok("a cold run stores exactly what a run with no cache stores", cold == want,
       sorted(set(cold) ^ set(want)))
    ok("so does a warm one, byte for byte", got == want,
       [k for k in want if got.get(k) != want.get(k)])
    cache_lines = [l for l in log_warm.splitlines() if l.startswith("[cache]")]
    ok("the warm run reused all three weekly reports", "reused 3 parse(s)" in log_warm,
       cache_lines)
    ok("...and parsed only the rent roll afresh, since a parse with names is never kept",
       "parsed 1 file(s) afresh" in log_warm, cache_lines)
    strip = lambda log: [l for l in log.splitlines() if not l.startswith("[cache]")]  # noqa: E731
    ok("every other log line is the one a run with no cache prints",
       strip(log_warm) == strip(log_off),
       [l for l in strip(log_warm) if l not in strip(log_off)][:3])

    build(warm, landed="2026-10-06T09:30:00Z")
    _, later = run(warm)
    stored = json.loads(later["the-landing/leasing_detail.json"])
    ok("a reused parse takes this run's arrival time, not the run that kept it",
       {w["landed_at"] for w in stored["weeks"]} == {"2026-10-06T09:30:00Z"},
       {w["landed_at"] for w in stored["weeks"]})

    held = b"".join(f.read_bytes() for f in (warm / "_cache").rglob("*") if f.is_file())
    ok("the leasing associate's name is nowhere in the cache", b"Edwin" not in held)
    ok("nor is the rent roll's resident", b"Resident Zed" not in held)
    ok("the daily reports' parses are what it holds",
       len(list((warm / "_cache" / "parse" / "entries").glob("*.pickle"))) == 3)


def _google_stubs():
    """fetch_drive imports the Drive client at module scope. Every call into
    it is replaced below, so where the library is absent a shell will do."""
    import types
    for name in ("google", "google.oauth2", "googleapiclient",
                 "googleapiclient.discovery", "googleapiclient.http"):
        sys.modules.setdefault(name, types.ModuleType(name))
    for mod, attr, val in (
            ("google.oauth2", "service_account", types.SimpleNamespace(
                Credentials=types.SimpleNamespace(
                    from_service_account_info=lambda *a, **k: None))),
            ("googleapiclient.discovery", "build", lambda *a, **k: None),
            ("googleapiclient.http", "MediaIoBaseDownload", object)):
        if not hasattr(sys.modules[mod], attr):
            setattr(sys.modules[mod], attr, val)


class Drive:
    """The Drive the fetch lists and downloads from: folders of files, each
    with the md5 Drive reports for it. Every download is recorded."""

    def __init__(self, fd):
        self.fd = fd
        self.blobs, self.tree = {}, {"R": []}
        self.downloads, self.fail = [], False
        fd._service = lambda: "svc"
        fd._list_children = self.list
        fd._download = self.download
        fd._build_svc = None

    def put(self, folder, name, data, fid=None):
        import hashlib
        fid = fid or f"{folder}/{name}"
        if not any(n == folder for _, n, _ in self.tree["R"]):
            self.tree["R"].append(("folder", folder, folder))
        self.tree.setdefault(folder, [])
        self.tree[folder] = [e for e in self.tree[folder] if e[2] != fid]
        self.tree[folder].append(("file", name, fid))
        self.blobs[fid] = (data, hashlib.md5(data).hexdigest())
        return fid

    def list(self, svc, pid, mime=None):
        out = []
        for kind, name, fid in self.tree.get(pid, []):
            f = {"id": fid, "name": name, "createdTime": "2026-10-05T12:00:00Z",
                 "modifiedTime": "2026-10-05T12:00:00Z",
                 "mimeType": ("application/vnd.google-apps.folder" if kind == "folder"
                              else "application/octet-stream")}
            if kind == "file":
                f.update(md5Checksum=self.blobs[fid][1], size=str(len(self.blobs[fid][0])))
            out.append(f)
        return out

    def download(self, svc, fid, dest):
        if self.fail == "partway":
            dest.parent.mkdir(parents=True, exist_ok=True)
            dest.write_bytes(self.blobs[fid][0][:10])
            raise ConnectionError("connection reset partway through")
        if self.fail:
            raise ConnectionError("Drive said 500")
        self.downloads.append(fid)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(self.blobs[fid][0])

    def took(self):
        got, self.downloads = self.downloads, []
        return got


REPORT_MAP = {"subfolders": [
    {"drive_folder": "Daily Leasing Reports", "report_type": "daily_leasing_report",
     "parser": "parse_daily_leasing", "status": "active", "file_glob": "*.xlsx"},
    {"drive_folder": "Rent Roll", "report_type": "rent_roll",
     "parser": "stub_roll_parser", "status": "active", "file_glob": "*.xlsx"},
    {"drive_folder": "EliseAI Reports", "report_type": "bldg_metrics_csv",
     "parser": "parse_building_metrics", "status": "active", "file_glob": "*.csv"},
]}


def section_fetch(tmp):
    print("\n7. the fetch leaves what it has parsed in Drive")
    _google_stubs()
    import fetch_drive as fd
    drive = Drive(fd)
    weeks = [("2026-09-08 Daily Report- Week Ending 9.7.26.xlsx", 5973, 3351),
             ("2026-09-15 Daily Report- Week Ending 9.14.26.xlsx", 6100, 5800),
             ("2026-09-22 Daily Report- Week Ending 9.21.26.xlsx", 5400, 5500)]
    src = pathlib.Path(tmp) / "fetch_src"
    src.mkdir()
    for name, rent, prior in weeks:
        build_week(src / name, rent, prior)
        drive.put("Daily Leasing Reports", name, (src / name).read_bytes())
    roll = {"report_type": "rent_roll", "property": "The Landing",
            "property_code": "p0005611", "as_of": "2026-09-21",
            "source_file": "RentRoll09_21_2026.xlsx", "totals": {"units": 1},
            "units": [{"unit": "901", "resident_name": "Resident Zed",
                       "rent": 5973, "occupied": True}], "checks": []}
    drive.put("Rent Roll", "RentRoll09_21_2026.xlsx", json.dumps(roll).encode())
    daily = {f"Daily Leasing Reports/{n}" for n, _, _ in weeks}

    def workdir(name):
        work = pathlib.Path(tmp) / name
        (work / "config").mkdir(parents=True)
        for f in pathlib.Path(REPO, "config").glob("*.json"):
            if f.name != "report_map.json":
                os.symlink(f, work / "config" / f.name)
        (work / "config" / "report_map.json").write_text(json.dumps(REPORT_MAP))
        return work

    @contextlib.contextmanager
    def inside(work, cache_env=None, key="sa-key"):
        saved = {k: os.environ.get(k) for k in ("PARSE_CACHE", "GDRIVE_SA_KEY",
                                                "GDRIVE_FOLDER_ID", "GITHUB_STEP_SUMMARY")}
        prev_cwd, prev_data, prev_argv = os.getcwd(), bm.DATA, sys.argv
        for k in saved:
            os.environ.pop(k, None)
        if cache_env is not None:
            os.environ["PARSE_CACHE"] = cache_env
        if key:
            os.environ["GDRIVE_SA_KEY"] = key
        os.environ["GDRIVE_FOLDER_ID"] = "R"
        bm.DATA = work / "data"
        os.chdir(work)
        try:
            yield
        finally:
            os.chdir(prev_cwd)
            bm.DATA, sys.argv = prev_data, prev_argv
            for k, v in saved.items():
                os.environ.pop(k, None)
                if v is not None:
                    os.environ[k] = v

    def fetch(work, *args, cache_env=None, fresh=True):
        if fresh:
            shutil.rmtree(work / "_downloads", ignore_errors=True)
        out = io.StringIO()
        with inside(work, cache_env), contextlib.redirect_stdout(out):
            sys.argv = ["fetch_drive.py", *args]
            fd.main()
        man = json.loads((work / "_downloads" / "manifest.json").read_text())
        return man, out.getvalue()

    def build(work, cache_env=None, key="sa-key"):
        shutil.rmtree(work / "data", ignore_errors=True)
        out = io.StringIO()
        with inside(work, cache_env, key), contextlib.redirect_stdout(out):
            bm.process_manifest()
        files = {str(p.relative_to(work / "data")): p.read_bytes()
                 for p in sorted((work / "data").rglob("*.json"))}
        return out.getvalue(), files

    def ids(man, left):
        return {f"{i['found_in']}/{i['name']}" for i in man
                if i.get("downloaded") is (not left)}

    entries = lambda work: list((work / "_cache" / "parse" / "entries").glob("*.pickle"))  # noqa: E731

    # The reference: every file downloaded and parsed, no cache at all.
    plain = workdir("fetch_plain")
    man, _ = fetch(plain, "--skip-parsed", cache_env="off")
    ok("with the cache off, --skip-parsed downloads every file",
       len(drive.took()) == 4 and all(i["downloaded"] for i in man))
    _, want = build(plain, cache_env="off")

    work = workdir("fetch_cached")
    man, _ = fetch(work, "--skip-parsed")
    ok("a cold fetch downloads every file", len(drive.took()) == 4)
    ok("...and records each one's Drive id and md5, the md5 of its bytes",
       all(i["drive_id"] and i["md5"] == drive.blobs[i["drive_id"]][1] for i in man))
    _, cold = build(work)
    ok("the cold build stores what the build with no cache stores", cold == want)

    man, flog = fetch(work, "--skip-parsed")
    ok("a warm fetch leaves the three kept weekly reports in Drive",
       ids(man, left=True) == daily, sorted(ids(man, left=True)))
    ok("...and downloads only the rent roll, whose parse is never kept",
       drive.took() == ["Rent Roll/RentRoll09_21_2026.xlsx"])
    ok("...and still lists all four in the manifest", len(man) == 4)
    ok("...each left one keyed on the md5 Drive reported",
       all(i["md5"] == drive.blobs[i["drive_id"]][1] for i in man if not i["downloaded"]))
    ok("the fetch says what it left and why",
       "3 left in Drive" in flog and "kept, not downloaded" in flog, flog[-300:])
    log, got = build(work)
    ok("the warm build stores exactly what the build with every file stores",
       got == want, [k for k in want if got.get(k) != want.get(k)])
    fetched = drive.took()
    ok("the build downloaded exactly one file: the one its self-check re-read",
       len(fetched) == 1 and fetched[0] in daily, fetched)
    ok("...and says so", log.count("[fetched]") == 1)
    ok("the build reused all three weekly reports' parses",
       "reused 3 parse(s)" in log, [l for l in log.splitlines() if "[cache]" in l])

    # main moved between the fetch and the build and the parses no longer
    # match -- here, every kept parse is gone. Every file is needed after all.
    man, _ = fetch(work, "--skip-parsed")
    drive.took()
    for e in entries(work):
        e.unlink()
    log, got = build(work)
    ok("a build whose kept parses do not answer downloads every file it needs",
       sorted(drive.took()) == sorted(daily))
    ok("...and still stores exactly what the build with every file stores", got == want)

    # The same, but Drive cannot be reached: the run must end, not carry on
    # without the files.
    man, _ = fetch(work, "--skip-parsed")
    for e in entries(work):
        e.unlink()
    drive.fail = True
    try:
        build(work)
        ended = None
    except SystemExit as e:
        ended = str(e)
    drive.fail = False
    ok("a file the build needs and cannot download ends the run",
       ended is not None and "could not be downloaded" in ended, ended)
    ok("...and no failure is kept for it, so the next run reads it",
       not entries(work))

    man, _ = fetch(work, "--skip-parsed")       # nothing kept: all four come down
    build(work)

    # Warm, but the one file the self-check re-reads cannot be downloaded: a
    # check that cannot run is not a check that failed.
    man, _ = fetch(work, "--skip-parsed")
    before = sorted(e.name for e in entries(work))
    drive.fail = True
    try:
        build(work)
        ended = None
    except SystemExit as e:
        ended = str(e)
    drive.fail = False
    ok("a self-check whose file cannot be downloaded ends the run",
       ended is not None and "could not be downloaded" in ended, ended)
    ok("...and sets no kept parse aside", sorted(e.name for e in entries(work)) == before)
    man, _ = fetch(work, "--skip-parsed")
    for e in entries(work):
        e.unlink()
    try:
        build(work, key=None)
        ended = None
    except SystemExit as e:
        ended = str(e)
    ok("without the Drive key in the step, the run ends naming the file",
       ended is not None and "GDRIVE_SA_KEY" in ended and "Daily Report" in ended, ended)

    # Back to a warm cache for the rest.
    fetch(work, "--skip-parsed")
    build(work)
    drive.took()

    man, _ = fetch(work)
    ok("without --skip-parsed every file is downloaded, cache or no cache",
       len(drive.took()) == 4 and all(i["downloaded"] for i in man))

    first = weeks[0][0]
    build_week(src / first, 6200, 3351)
    drive.put("Daily Leasing Reports", first, (src / first).read_bytes())
    man, _ = fetch(work, "--skip-parsed")
    ok("a report whose bytes changed in Drive is downloaded again",
       f"Daily Leasing Reports/{first}" in drive.took())
    build(work)

    # A copy on disk from an earlier local run, with other bytes than Drive's.
    man, _ = fetch(work, "--skip-parsed")
    drive.took()
    stale = work / "_downloads" / "daily_leasing_report" / weeks[1][0]
    stale.parent.mkdir(parents=True, exist_ok=True)
    stale.write_bytes(b"an older copy")
    man, _ = fetch(work, "--skip-parsed", fresh=False)
    took = drive.took()
    ok("a stale local copy of a file left in Drive is removed, so nothing reads it",
       not stale.exists() and not set(took) & daily, (stale.exists(), took))
    same = work / "_downloads" / "daily_leasing_report" / weeks[2][0]
    same.write_bytes((src / weeks[2][0]).read_bytes())
    fetch(work, "--skip-parsed", fresh=False)
    ok("...while a local copy that IS Drive's file is left where it is", same.exists())

    # The same filename in two folders: the second must not be taken, whether
    # or not the first was downloaded.
    drive.took()
    drive.put("Leasing Copies", weeks[1][0], b"same name, other folder", fid="dup")
    REPORT_MAP["subfolders"].append(
        {"drive_folder": "Leasing Copies", "report_type": "daily_leasing_report",
         "parser": "parse_daily_leasing", "status": "active", "file_glob": "*.xlsx"})
    (work / "config" / "report_map.json").write_text(json.dumps(REPORT_MAP))
    man, flog = fetch(work, "--skip-parsed")
    REPORT_MAP["subfolders"].pop()
    (work / "config" / "report_map.json").write_text(json.dumps(REPORT_MAP))
    paths = [i["path"] for i in man]
    ok("a second file of the same name is not taken over one left in Drive",
       len(paths) == len(set(paths)) and "dup" not in drive.took()
       and "already taken" in flog, flog[-400:])
    drive.tree["R"] = [e for e in drive.tree["R"] if e[1] != "Leasing Copies"]

    # A type something other than the parser reads from disk is always fetched.
    drive.put("EliseAI Reports", "metrics-building-2026-10-05.csv", b"a,b\n1,2\n")
    probe = {"path": "_downloads/bldg_metrics_csv/metrics-building-2026-10-05.csv",
             "md5": drive.blobs["EliseAI Reports/metrics-building-2026-10-05.csv"][1]}
    with inside(work):
        cache = pc.ParseCache(pc.DEFAULT_DIR)
        k = cache.key(probe, pc._Named("parse_building_metrics",
                                        pathlib.Path(HERE) / "parse_building_metrics.py"))
        (cache.entries / f"{k}.pickle").write_bytes(b"")
        held = cache.has(probe, "parse_building_metrics")
    drive.took()
    man, _ = fetch(work, "--skip-parsed")
    ok("a building-metrics CSV is downloaded even when its parse is kept, "
       "since the scorecard fill reads it from disk",
       held and "EliseAI Reports/metrics-building-2026-10-05.csv" in drive.took())

    # A fetch hook that raises an ordinary exception: the file's parse is
    # unknown, not failed, so the run ends and nothing is kept for it.
    with inside(work):
        it = dict(next(i for i in man if not i["downloaded"]))
        pathlib.Path(it["path"]).unlink(missing_ok=True)
        cache = pc.ParseCache(pc.DEFAULT_DIR, log=lambda *_: None,
                              fetch=lambda i: (_ for _ in ()).throw(RuntimeError("no route")))
        mod = importlib.import_module(it["parser"])
        k = cache.key(it, mod)
        (cache.entries / f"{k}.pickle").unlink(missing_ok=True)
        try:
            cache.parse(it, mod)
            raised = None
        except BaseException as e:                      # noqa: BLE001
            raised = e
        kept_after = (cache.entries / f"{k}.pickle").exists()
    ok("a hook that fails with an ordinary exception still ends the run, and "
       "nothing is kept as the file's failure",
       isinstance(raised, SystemExit) and "no route" in str(raised)
       and not kept_after, (raised, kept_after))

    # fetch_one on its own: a download cut off partway leaves nothing behind
    # for a later step -- or a later local run -- to read as the file.
    item = next(i for i in man if not i["downloaded"])
    drive.fail = "partway"
    with inside(work):
        pathlib.Path(item["path"]).unlink(missing_ok=True)
        try:
            fd.fetch_one(item)
            ended = None
        except SystemExit as e:
            ended = str(e)
        partial = pathlib.Path(item["path"]).exists()
    drive.fail = False
    ok("a download cut off partway ends the run and leaves no partial file",
       ended is not None and "could not be downloaded" in ended and not partial,
       (ended, partial))

    # fetch_one on its own: a file that changed between the listing and the
    # download is not parsed under the old copy's key.
    item = next(i for i in man if not i["downloaded"])
    item = dict(item, md5="0" * 32)
    with inside(work):
        try:
            fd.fetch_one(item)
            ended = None
        except SystemExit as e:
            ended = str(e)
        gone = not pathlib.Path(item["path"]).exists()
    ok("a file whose bytes no longer match the listing ends the run, and is not left on disk",
       ended is not None and "changed in Drive" in ended and gone, ended)


def section_contract():
    print("\n8. what the fetch may leave in Drive, and the workflow around it")
    import re
    sys.path.insert(0, HERE)
    import fetch_drive as fd
    rtypes = {e.get("report_type") for e in json.load(
        open(os.path.join(REPO, "config", "report_map.json")))["subfolders"]}
    # Everything the daily run executes, and every workflow that skips.
    skipping = [p for p in pathlib.Path(REPO, ".github", "workflows").glob("*.yml")
                if "--skip-parsed" in p.read_text()]
    sources = [p for p in pathlib.Path(HERE).iterdir()
               if p.suffix in (".py", ".sh") and not p.name.startswith("test_")
               and p.name != "fetch_drive.py"] + skipping
    readers = {}
    for p in sources:
        for m in re.finditer(r"_downloads/([A-Za-z_]+)", p.read_text()):
            if m.group(1) in rtypes:
                readers.setdefault(m.group(1), set()).add(p.name)
    stray = {t: sorted(f) for t, f in readers.items() if t not in fd.READ_FROM_DISK}
    ok("every report type read from _downloads/ outside the parse is one the fetch "
       "always downloads (READ_FROM_DISK)", not stray, stray)
    ok("...and the scorecard fill's CSV is one of them",
       "bldg_metrics_csv" in readers and "bldg_metrics_csv" in fd.READ_FROM_DISK, readers)

    wf = pathlib.Path(REPO, ".github", "workflows", "update.yml").read_text()
    steps = {m.group(1).strip(): m.start()
             for m in re.finditer(r"^      - name: (.+)$", wf, re.M)}
    body = lambda name: wf[steps[name]:min([v for v in steps.values()  # noqa: E731
                                            if v > steps[name]] + [len(wf)])]
    ok("update.yml restores the parse cache before the fetch reads it",
       steps["Restore the parse cache"] < steps["Fetch reports from Google Drive"])
    ok("...and the fetch is told to skip what it has parsed",
       "--skip-parsed" in body("Fetch reports from Google Drive"))
    ok("...and both steps that build can download what the fetch left in Drive",
       all("GDRIVE_SA_KEY" in body(n) for n in
           ("Build metrics, fill the scorecard, build the lineage",
            "Check for personal data, commit and push")))
    others = [p.name for p in pathlib.Path(REPO, ".github", "workflows").glob("*.yml")
              if p.name != "update.yml" and "fetch_drive.py" in p.read_text()
              and p in skipping]
    ok("no other workflow that fetches skips, since none has the build's cache",
       not others, others)


def main():
    tmp = tempfile.mkdtemp(prefix="parse_cache_")
    try:
        sb = Sandbox(tmp)
        section_key(sb)
        section_kept(sb)
        section_check(sb)
        section_housekeeping(sb)
        section_real_keys()
        section_end_to_end(tmp)
        section_fetch(tmp)
        section_contract()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"\n{PASS} passed, {FAIL} failed")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
