#!/usr/bin/env python3
"""Guard tests for the occupancy split behind the Unit Inventory card.

The card draws leased against vacant per bedroom type, from two reports that
have never met: the rent roll knows which units are let, the unit directory
knows what a floorplan is. Three things can go wrong quietly, so all three are
checked here:

  * the wrong units counted as leased -- a unit on notice or holding over is
    still occupied, and a vacant unit can carry a market rent;
  * the two sources counting the building differently, which is guaranteed
    unless the directory's waitlist placeholders are excluded -- they are not
    apartments and the rent roll has never heard of them;
  * anything unit level reaching the published block. The roll arrives with
    resident names and is gitignored for that reason; only counts per floorplan
    ride out in metrics.json, and that is what makes the card possible at all.

The property line is checked too, because a roll the pipeline cannot attribute
is skipped whole and the card silently keeps its unsplit bars.

Fixture-free and offline: both exports are built in a temp dir.

Run: python scripts/test_occupancy.py
"""
import os
import sys
import json
import shutil
import pathlib
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_metrics as bm        # noqa: E402
import parse_rent_roll as rr      # noqa: E402
import parse_unit_directory as ud  # noqa: E402

PASS = FAIL = 0


def ok(name, cond, detail=None):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}" + (f"\n        {detail}" if detail is not None else ""))


# (unit, plan, sqft, market, resident, actual, move_out)
# Two 1-beds let, one 1-bed empty; both 2-beds let, one of them on notice.
# 402 is the trap: a vacant unit still carries a market rent and a square
# footage, so "has numbers in it" is not the same as "has someone in it".
ROLL = [
    ("101", "laa1",  560, 5595, "t0001", 5400, None),
    ("102", "laa1",  536, 5640, "t0002", 5500, None),
    ("402", "laa1",  540, 5610, None,       0, None),
    ("355", "lab19", 1095, 8241, "t0003", 8000, None),
    ("356", "lab19", 1090, 8200, "t0004", 8100, "2026-10-31"),
]
UD_HEADER = ["Unit", "Address", "Unit Type", "Rent", "Deposit", "Sqft",
             "Room", "Baths", "Notes"]
# The directory lists the same five apartments plus a waitlist placeholder,
# which is what makes the two sources disagree until it is excluded.
UD_ROWS = [
    ("101", "laa1", 5595, 560, 1, 1),
    ("102", "laa1", 5640, 536, 1, 1),
    ("402", "laa1", 5610, 540, 1, 1),
    ("355", "lab19", 8241, 1095, 2, 2),
    ("356", "lab19", 8200, 1090, 2, 2),
    ("WAITLIST", "laa1", 5565, 821, 1, 1),
]


def build_roll(path, rows=ROLL, applicant=True, named_at_bottom=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rent Roll"
    ws.append(["Rent Roll"])
    ws.append(["For Selected Properties"])
    ws.append(["As of 07/14/2026"])
    ws.append(["Unit", "Unit Type", "Sq Ft", "Resident", "Name", "Market Rent",
               "Actual Rent", "Resident Deposit", "Other Deposit", "Move In",
               "Lease Expiration", "Move Out", "Balance"])
    for unit, plan, sqft, mkt, res, act, out in rows:
        ws.append([unit, plan, sqft, res, ("Resident " + unit) if res else None,
                   mkt, act, 0, 0, None, None, out, 0])
    ws.append(["Total", None, None, None, None,
               sum(r[3] for r in rows), sum(r[5] for r in rows), 0, 0,
               None, None, None, 0])
    # The applicant section below the total is the other trap the parser stops
    # at; a roll that runs straight on would count applicants as units.
    if applicant:
        ws.append(["Future Residents/Applicants"])
        ws.append(["901", "laa1", 560, "a0001", "Applicant", 5600, 5600, 0, 0,
                   None, None, None, 0])
    if named_at_bottom:
        # What RentRoll09_11_2026 actually does: the header says only "For
        # Selected Properties" and the buildings are named in a summary block
        # below every unit row.
        ws.append([None])
        ws.append([None, None, None, None, "The Landing(p0005611)"])
        ws.append([None, None, None, None, "The Landing - PDR(p0005640)"])
    wb.save(path)


def build_dir(path, rows=UD_ROWS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report1"
    ws.append(["Unit Directory"])
    ws.append(["For Selected Properties"])
    ws.append(UD_HEADER)
    ws.append(["p0005611"])
    for unit, plan, rent, sf, beds, baths in rows:
        ws.append([unit, "somewhere", plan, float(rent), 0.0, float(sf),
                   beds, float(baths), None])
    ws.append([f"Total p0005611", None, f"Units: {len(rows)}",
               float(sum(r[2] for r in rows)), 0.0, float(sum(r[3] for r in rows))])
    ws.append(["Grand Total", None, f"Units: {len(rows)}",
               float(sum(r[2] for r in rows)), 0.0, float(sum(r[3] for r in rows))])
    wb.save(path)


def bed_rollup(by_plan, plans, residential=False):
    """What the card does: the roll's plan counts rolled onto bedroom groups,
    beside the directory's own census of the same groups."""
    out = {}
    key = "units"
    census = {}
    for code, p in plans.items():
        k = "?" if p.get("bedrooms") is None else p["bedrooms"]
        census[k] = census.get(k, 0) + (p.get(key) or 0)
    for code, o in by_plan.items():
        p = plans.get(code)
        k = "?" if (not p or p.get("bedrooms") is None) else p["bedrooms"]
        t = out.setdefault(k, {"units": 0, "leased": 0, "vacant": 0})
        for f in ("units", "leased", "vacant"):
            t[f] += o.get(f, 0)
    return census, out


def main():
    tmp = tempfile.mkdtemp()
    try:
        roll_p = os.path.join(tmp, "rent_roll.xlsx")
        dir_p = os.path.join(tmp, "UnitDirectory07_14_2026.xlsx")
        build_roll(roll_p)
        build_dir(dir_p)
        roll = rr.parse(roll_p)
        drc = ud.parse(dir_p)
        plans = drc["sections"][0]["plans"]

        print("what counts as leased")
        summ = bm.rent_roll_summary(roll)
        occ = {"by_plan": summ["by_plan"], "units": summ["units"],
               "leased": summ["occupied"], "vacant": summ["vacant"]}
        ok("five units parsed, applicants excluded", occ["units"] == 5, occ)
        ok("four leased, one vacant",
           occ["leased"] == 4 and occ["vacant"] == 1, occ)
        ok("a unit on notice is still leased",
           occ["by_plan"]["lab19"]["leased"] == 2, occ["by_plan"])
        ok("a vacant unit carrying a market rent is not leased",
           occ["by_plan"]["laa1"]["vacant"] == 1, occ["by_plan"])
        ok("leased + vacant accounts for every unit, per plan",
           all(p["leased"] + p["vacant"] == p["units"]
               for p in occ["by_plan"].values()), occ["by_plan"])
        ok("...and in total",
           occ["leased"] + occ["vacant"] == occ["units"], occ)

        print("\nnothing unit level is published")
        prev = bm.DATA
        bm.DATA = pathlib.Path(tmp) / "data"
        try:
            bm.store_rent_roll({"name": "Fixture", "slug": "fx"}, roll)
        finally:
            bm.DATA = prev
        stored = json.load(open(pathlib.Path(tmp) / "data" / "fx" / "rent_roll.json"))
        pub = bm.rent_roll_summary(stored)["by_plan"]
        blob = json.dumps(pub)
        ok("by_plan survives the store and the scrub", bool(pub), pub)
        ok("no resident name anywhere in it", "Resident 101" not in blob, blob[:200])
        ok("no resident code either", "t0001" not in blob, blob[:200])
        ok("no unit id either", '"101"' not in blob, blob[:200])
        ok("counts only, nothing else",
           all(set(v) == {"units", "leased", "vacant"} for v in pub.values()), pub)
        ok("the per-unit rows stay in the gitignored file",
           (pathlib.Path(tmp) / "data" / "fx" / "rent_roll.json").exists())
        ok("and the scrub dropped names from that file too",
           "resident_name" not in json.dumps(stored)
           and "Resident 101" not in json.dumps(stored),
           "a person-shaped field survived the scrub")

        print("\nthe roll is its own census; the directory only supplies bedrooms")
        # The two count different things on purpose -- the directory lists the
        # waitlist placeholder, the roll lists leasable apartments -- so the
        # card follows the roll and reports the difference rather than
        # reconciling them into one number that is neither.
        ok("the directory counts the placeholder, the roll does not",
           drc["sections"][0]["units"] == 6
           and drc["sections"][0]["residential_units"] == 5,
           drc["sections"][0])
        raw_census, rolled = bed_rollup(occ["by_plan"], plans, residential=False)
        ok("the directory's census over-counts the 1-bed group, and that is fine",
           raw_census[1] == 4 and rolled[1]["units"] == 3, (raw_census, rolled))
        ok("the difference is visible, so the note can report it rather than "
           "reconcile it",
           sum(raw_census.values()) - sum(g["units"] for g in rolled.values()) == 1,
           (raw_census, rolled))
        ok("the rollup is the split the bars draw",
           rolled[1] == {"units": 3, "leased": 2, "vacant": 1}
           and rolled[2] == {"units": 2, "leased": 2, "vacant": 0}, rolled)
        ok("every bar is exactly its two segments, whatever the directory says",
           all(g["leased"] + g["vacant"] == g["units"] for g in rolled.values()),
           rolled)

        print("\na plan the directory has never heard of is not silently dropped")
        roll2_p = os.path.join(tmp, "roll2.xlsx")
        build_roll(roll2_p, ROLL + [("777", "zzz9", 700, 6000, "t0009", 5900, None)])
        s2 = bm.rent_roll_summary(rr.parse(roll2_p))
        occ2 = {"by_plan": s2["by_plan"], "units": s2["units"]}
        census2, rolled2 = bed_rollup(occ2["by_plan"], plans)
        ok("it lands in the unknown-bedroom group rather than vanishing",
           rolled2.get("?", {}).get("units") == 1, rolled2)
        ok("so leased + vacant still accounts for every unit on the roll",
           sum(g["units"] for g in rolled2.values()) == occ2["units"],
           (rolled2, occ2["units"]))
        ok("and that group is still internally consistent",
           all(g["leased"] + g["vacant"] == g["units"] for g in rolled2.values()),
           rolled2)
        ok("the plan is nameable, so the note can say which one it is",
           [c for c in occ2["by_plan"] if c not in plans] == ["zzz9"],
           list(occ2["by_plan"]))

        print("\nthe property line is found wherever Yardi put it")
        # Without this the whole roll is skipped as unattributed and the card
        # never sees a split at all -- which is what was happening.
        bottom_p = os.path.join(tmp, "bottom.xlsx")
        build_roll(bottom_p, named_at_bottom=True)
        import openpyxl as _ox
        ws = rr._pick_sheet(_ox.load_workbook(bottom_p, data_only=True))
        ok("a property named below the unit rows is still found",
           rr._property(ws) == ("The Landing", "p0005611"), rr._property(ws))
        ok("the residential code wins over the commercial one beside it",
           rr._property(ws)[1] == "p0005611", rr._property(ws))
        ok("a roll naming nothing at all still says so rather than guessing",
           rr._property(rr._pick_sheet(_ox.load_workbook(roll_p, data_only=True)))
           == (None, None))
        ok("and the bottom block is not read as a unit row",
           rr.parse(bottom_p)["totals"]["units"] == 5,
           rr.parse(bottom_p)["totals"]["units"])

        print("\na unit row with no plan code still balances")
        roll3_p = os.path.join(tmp, "roll3.xlsx")
        build_roll(roll3_p, ROLL + [("888", None, 700, 6000, None, 0, None)])
        s3 = bm.rent_roll_summary(rr.parse(roll3_p))
        occ3 = {"by_plan": s3["by_plan"], "units": s3["units"],
                "leased": s3["occupied"], "vacant": s3["vacant"]}
        ok("it is counted under a named bucket, not dropped",
           occ3["by_plan"].get("(no plan)", {}).get("units") == 1,
           occ3["by_plan"])
        ok("totals still balance", occ3["leased"] + occ3["vacant"] == occ3["units"],
           occ3)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print(f"\n{PASS} passed, {FAIL} failed")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
