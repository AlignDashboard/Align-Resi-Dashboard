#!/usr/bin/env python3
"""Guard tests: the parsers read a sheet's dimensions once, not per row.

openpyxl does not store `max_row` or `max_column`. Each access scans every cell
in the sheet. That is free on a 300-row report and ruinous on a sheet formatted
out to tens of thousands of rows: the Madelon daily report is 51 MB of XML and
about 2.3 million cells, and `xlsx_anchors.header_map` asked for `max_row` once
per cell it read -- 3,338 full scans, 277 of the 323 seconds that one parse
took. There is one such file per day and every run re-parses them all, so the
daily build grew by minutes a day until GitHub's six-hour limit cancelled it
(2026-09-24 onward, runs #94-#97).

The fix is `xlsx_anchors.dims(ws)`, read once before a loop. Nothing here times
anything -- a timing test would pass on a fast machine and flake on a slow one.
Instead each helper runs against a sheet of a few hundred rows with openpyxl's
two properties wrapped in a counter, and must read them a constant number of
times however long the sheet is. Output is checked too: the counter must not
change what the helper returns.

Fixture-free: the workbooks are built in memory.

Run: python scripts/test_xlsx_dims.py
"""

import os
import sys

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import parse_daily_leasing  # noqa: E402
import parse_delinquency  # noqa: E402
import parse_rent_roll  # noqa: E402
import parse_unit_directory  # noqa: E402
import xlsx_anchors  # noqa: E402

PASS = FAIL = 0


def check(name, got, want):
    global PASS, FAIL
    if got == want:
        PASS += 1
        print(f"   PASS {name}")
    else:
        FAIL += 1
        print(f"   FAIL {name}\n        got  {got!r}\n        want {want!r}")


READS = {"n": 0}
_ROW, _COL = Worksheet.max_row, Worksheet.max_column


def _counted(prop):
    def get(self):
        READS["n"] += 1
        return prop.fget(self)
    return property(get)


Worksheet.max_row = _counted(_ROW)
Worksheet.max_column = _counted(_COL)

# However long the sheet, a helper may read each dimension a handful of times.
# The old code read them per row or per cell: hundreds to thousands here.
BOUND = 4


def reads(fn, *args):
    READS["n"] = 0
    out = fn(*args)
    return out, READS["n"]


def sheet(rows, header, filler=True):
    """A sheet with `header` on its own row and `rows` rows of text beneath."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, line in header:
        for c, v in enumerate(line, 1):
            ws.cell(row=r, column=c, value=v)
    if filler:
        for r in range(max(h[0] for h in header) + 1, rows + 1):
            for c in range(1, 21):
                ws.cell(row=r, column=c, value=f"r{r}c{c}")
    return ws


for n in (200, 800):
    print(f"sheets of {n} rows x 20 columns")

    ws = sheet(n, [(3, ["Unit", "Market Rent", "Actual Rent", "Resident"])])
    spec = {"unit": r"^unit$", "market": r"market", "actual": r"actual"}
    (fields, hdr), k = reads(xlsx_anchors.header_map, ws, spec)
    check(f"{n}: header_map finds the header", (fields, hdr),
          ({"unit": 1, "market": 2, "actual": 3}, 3))
    check(f"{n}: header_map reads the dimensions at most {BOUND} times",
          k <= BOUND, True)

    ws = sheet(n, [(2, ["Property:", "Chorus"]), (3, ["Units:", 416]),
                   (5, ["Week Ending", "2026-09-20"])])
    got, k = reads(parse_daily_leasing._sheet_property, ws)
    check(f"{n}: daily leasing reads its property block", got, ("Chorus", 416))
    check(f"{n}: ...reading the dimensions at most {BOUND} times", k <= BOUND, True)
    got, k = reads(parse_daily_leasing._week_ending, ws, "x.xlsx")
    check(f"{n}: daily leasing reads its week-ending cell", got, "2026-09-20")
    check(f"{n}: ...reading the dimensions at most {BOUND} times", k <= BOUND, True)

    # deep in the 40-row search, so a per-row read has rows to repeat on
    ws = sheet(n, [(30, ["Delinquency Summary as of 14 September 2026"])])
    got, k = reads(parse_delinquency._as_of, ws)
    check(f"{n}: delinquency reads its as-of line", got, "2026-09-14")
    check(f"{n}: ...reading the dimensions at most {BOUND} times", k <= BOUND, True)

    ws = sheet(n, [(12, ["Unit", "Unit Type", "SqFt", "Rooms"])])
    (got, row), k = reads(parse_unit_directory._header, ws)
    check(f"{n}: unit directory finds its header", (sorted(got), row),
          (["beds", "plan", "sqft", "unit"], 12))
    check(f"{n}: ...reading the dimensions at most {BOUND} times", k <= BOUND, True)

    ws = sheet(n, [(1, ["Rent Roll as of = 09/21/2026"])])
    got, k = reads(parse_rent_roll._text_cells, ws, n, None)
    check(f"{n}: rent roll's text scan sees every text cell", len(got), 1 + (n - 1) * 20)
    check(f"{n}: ...reading the dimensions at most {BOUND} times", k <= BOUND, True)
    ws.cell(row=n, column=1, value="Total")
    ws.cell(row=n, column=2, value=1000.0)
    ws.cell(row=n, column=3, value=900.0)
    _, k = reads(parse_rent_roll._report_totals, ws,
                 {"unit": 1, "market_rent": 2, "actual_rent": 3})
    check(f"{n}: rent roll's total-row search reads the dimensions at most "
          f"{BOUND} times", k <= BOUND, True)

Worksheet.max_row, Worksheet.max_column = _ROW, _COL
print(f"\n{PASS} passed, {FAIL} failed")
sys.exit(1 if FAIL else 0)
